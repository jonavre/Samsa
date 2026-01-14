from flask import Flask, request, jsonify, render_template, send_file, session, redirect
from flask_cors import CORS
from functools import wraps
import psycopg2
from psycopg2.extras import RealDictCursor
from dateutil import parser
import datetime
import io
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from flask import request, flash
from werkzeug.utils import secure_filename
import os
from flask import Flask, request, redirect, url_for, render_template, session, send_file
from io import BytesIO
from flask import Response

# ================================
# CONFIGURACIÓN BD (RAILWAY)
# ================================
DB_HOST = "switchback.proxy.rlwy.net"
DB_PORT = 38554
DB_NAME = "railway"
DB_USER = "postgres"
DB_PASSWORD = "pZfEFTAfgrMrWWWlLovDGFaIfLMAjtIt"

app = Flask(__name__, template_folder="templates")
app.secret_key = "samsa_secreta"
CORS(app)

def get_connection():
    try:
        return psycopg2.connect(
            host=DB_HOST,
            port=DB_PORT,
            dbname=DB_NAME,
            user=DB_USER,
            password=DB_PASSWORD,
            sslmode="require"
        )
    except Exception as e:
        print("❌ Error conexión BD:", e)
        return None

# ================================
# DECORADOR DE ROLES
# ================================
def role_required(*roles):
    def wrapper(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if "rol" not in session or session["rol"] not in roles:
                return redirect("/")
            return f(*args, **kwargs)
        return decorated
    return wrapper

# ================================
# LOGIN
# ================================
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        nombre = request.form.get("nombre")
        password = request.form.get("password")

        conn = get_connection()
        if not conn:
            return render_template("login.html", error="Error de conexión")

        cur = conn.cursor()
        cur.execute("""
            SELECT nombre, rol
            FROM credenciales
            WHERE nombre=%s AND contrasena=%s
        """, (nombre, password))

        user = cur.fetchone()
        cur.close()
        conn.close()

        if user:
            session["usuario"] = user[0]
            session["rol"] = user[1]
            return redirect("/")   
        else:
            return render_template("login.html", error="Credenciales incorrectas")

    return render_template("login.html")

# ================================
# HOME SEGÚN ROL
# ================================
@app.route("/")
def home():
    if "rol" not in session:
        return redirect("/login")

    conn = get_connection()
    cur = conn.cursor()
    cur.execute("SELECT id FROM programacion ORDER BY id DESC LIMIT 1")
    fila = cur.fetchone()
    ultima_programacion_id = fila[0] if fila else None
    cur.close()
    conn.close()

    return render_template("index.html", ultima_programacion_id=ultima_programacion_id)

@app.route("/logout")
def logout():
    session.clear()
    return redirect("/login")

# ================================
# FRONTEND (PROTEGIDO)
# ================================
@app.route("/produccion")
@role_required("administrador")
def produccion():
    return render_template("produccion.html")

@app.route("/producciones")
@role_required("administrador")
def producciones():
    return render_template("producciones.html")

@app.route("/lotes")
@role_required("administrador")
def lotes():
    return render_template("lotes.html")

@app.route("/reportes")
@role_required("administrador")
def reportes():
    return render_template("reportes.html")

@app.route("/config")
@role_required("administrador")
def config():
    return render_template("config.html")

@app.route("/pedidos")
@role_required("administrador", "operario")
def pedidos():
    return render_template("pedidos.html")

@app.route("/afalpi")
def afalpi():
    return render_template("afalpi.html")

@app.route("/inventario")
def inventario():
    return render_template("inventario.html")

@app.route("/solicitudes")
@role_required("administrador")
def solicitudes():
    return render_template("solicitudes.html")

# ================================
# STATUS API
# ================================
@app.route("/api/status")
def status():
    return jsonify({"status": "API SAMSA funcionando correctamente"})

# ================================
# SUBIR FOTO
# ================================
UPLOAD_FOLDER = os.path.join('static', 'uploads')
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@app.route("/programacion/<int:id>")
def mostrar_programacion(id):
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("SELECT nombre_archivo, contenido FROM programacion WHERE id=%s", (id,))
    fila = cur.fetchone()
    cur.close()
    conn.close()

    if fila:
        nombre_archivo, contenido = fila
        return Response(contenido, mimetype="image/png")
    return "Imagen no encontrada", 404

# ================================
# LISTAR PRODUCCIONES
# ================================
@app.route("/api/producciones", methods=["GET"])
def obtener_producciones():
    conn = get_connection()
    if not conn:
        return jsonify([]), 500

    cur = conn.cursor()
    try:
        cur.execute("""
            SELECT id, fruta, fecha, tanda, codigo_lote, proceso,
                   presentacion, cantidad, kilos_totales, ph, brix,
                   azucar_kg, pectina_kg, observaciones
            FROM produccion
            ORDER BY id
        """)
        rows = cur.fetchall()
        data = []
        for r in rows:
            data.append({
                "id": r[0],
                "fruta": r[1],
                "fecha": r[2].isoformat() if r[2] else "",
                "tanda": r[3],
                "codigo_lote": r[4],
                "proceso": r[5],
                "presentacion": r[6],
                "cantidad": float(r[7]) if r[7] is not None else 0,
                "kilos_totales": float(r[8]) if r[8] is not None else 0,
                "ph": float(r[9]) if r[9] is not None else 0,
                "brix": float(r[10]) if r[10] is not None else 0,
                "azucar_kg": float(r[11]) if r[11] is not None else 0,
                "pectina_kg": float(r[12]) if r[12] is not None else 0,
                "observaciones": r[13]
            })
        return jsonify(data)
    except Exception as e:
        print("❌ Error al obtener producciones:", e)
        return jsonify([]), 500
    finally:
        cur.close()
        conn.close()

# ================================
# BULK INSERT / UPDATE
# ================================
@app.route("/api/producciones/bulk", methods=["POST"])
def guardar_bulk():
    data = request.json
    if not data:
        return jsonify({"error": "No hay datos para guardar"}), 400

    conn = get_connection()
    if not conn:
        return jsonify({"error": "No se pudo conectar a la BD"}), 500

    cur = conn.cursor()
    try:
        for fila in data:
            if not fila.get("fruta") and not fila.get("codigo_lote"):
                continue

            fecha = None
            if fila.get("fecha"):
                try:
                    fecha = parser.parse(fila["fecha"]).date()
                except:
                    fecha = None

            cantidad = fila.get("cantidad") or 0
            kilos_totales = fila.get("kilos_totales") or 0
            id_val = fila.get("id")

            if id_val:
                cur.execute("""
                    UPDATE produccion SET
                        fruta=%s, fecha=%s, tanda=%s, codigo_lote=%s,
                        proceso=%s, presentacion=%s, cantidad=%s,
                        unidad='kg', kilos_totales=%s, ph=%s, brix=%s,
                        azucar_kg=%s, pectina_kg=%s, observaciones=%s
                    WHERE id=%s
                """, (
                    fila.get("fruta"), fecha, fila.get("tanda"),
                    fila.get("codigo_lote"), fila.get("proceso"),
                    fila.get("presentacion"), cantidad,
                    kilos_totales, fila.get("ph"), fila.get("brix"),
                    fila.get("azucar_kg"), fila.get("pectina_kg"),
                    fila.get("observaciones"), int(id_val)
                ))
            else:
                cur.execute("""
                    INSERT INTO produccion (
                        fruta, fecha, tanda, codigo_lote, proceso,
                        presentacion, cantidad, unidad, kilos_totales,
                        ph, brix, azucar_kg, pectina_kg, observaciones
                    ) VALUES (%s,%s,%s,%s,%s,%s,%s,'kg',%s,%s,%s,%s,%s,%s)
                """, (
                    fila.get("fruta"), fecha, fila.get("tanda"),
                    fila.get("codigo_lote"), fila.get("proceso"),
                    fila.get("presentacion"), cantidad,
                    kilos_totales, fila.get("ph"), fila.get("brix"),
                    fila.get("azucar_kg"), fila.get("pectina_kg"),
                    fila.get("observaciones")
                ))

        conn.commit()
        return jsonify({"mensaje": "OK"})
    except Exception as e:
        conn.rollback()
        print("❌ Error en bulk:", e)
        return jsonify({"error": str(e)}), 500
    finally:
        cur.close()
        conn.close()


# ================================
# CALENDARIO
# ================================
@app.route("/api/calendario/<fecha>")
def obtener_codigo_calendario(fecha):
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("SELECT codigo FROM calendario_produccion WHERE fecha=%s", (fecha,))
    row = cur.fetchone()
    cur.close()
    conn.close()
    return jsonify({"codigo_dia": row[0] if row else None})

# ================================
# PRODUCTOS
# ================================
@app.route("/api/productos")
def obtener_productos():
    tipo = request.args.get("tipo")
    conn = get_connection()
    cur = conn.cursor(cursor_factory=RealDictCursor)
    if tipo:
        cur.execute("SELECT codigo, descripcion FROM productos WHERE tipo=%s", (tipo,))
    else:
        cur.execute("SELECT codigo, descripcion FROM productos")
    data = cur.fetchall()
    cur.close()
    conn.close()
    return jsonify(data)

# ================================
# EXPORTAR EXCEL
# ================================
@app.get("/exportar_excel")
def exportar_excel():
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("""
        SELECT fecha, fruta, proceso, codigo_lote, ph, brix,
               presentacion, kilos_totales, azucar_kg, pectina_kg, observaciones
        FROM produccion
        ORDER BY fecha DESC
    """)
    rows = cur.fetchall()
    cur.close()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Producciones"
    ws.append([
        "FECHA","FRUTA","PROCESO","LOTE","PH","BRIX",
        "PRESENTACIÓN","CANTIDAD PROCESADA (KG)",
        "AZÚCAR (KG)","PECTINA","OBSERVACIONES"
    ])
    for r in rows:
        ws.append(list(r))

    table = Table(displayName="ProduccionTabla", ref=f"A1:K{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(table)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="producciones.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route("/api/afalpi/<int:id>", methods=["DELETE"])
@role_required("administrador")
def borrar_afalpi(id):
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("DELETE FROM afalpi WHERE id=%s", (id,))
        conn.commit()
        return jsonify({"ok": True})
    except Exception as e:
        conn.rollback()
        return jsonify({"error": str(e)}), 500
    finally:
        cur.close()
        conn.close()

@app.route("/api/inventario/<int:id>", methods=["DELETE"])
@role_required("administrador")
def borrar_inventario(id):
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute("DELETE FROM inventario WHERE id=%s", (id,))
        conn.commit()
        return jsonify({"ok": True})
    except Exception as e:
        conn.rollback()
        return jsonify({"error": str(e)}), 500
    finally:
        cur.close()
        conn.close()

# ================================
# CREAR PEDIDO AFALPI
# ================================

@app.route("/api/pedidos/afalpi", methods=["POST"])
def crear_pedido_afalpi():
    try:
        data = request.get_json()
        comentario = data.get("comentario")

        if not comentario:
            return jsonify({"error": "Comentario requerido"}), 400

        conn = get_connection()

        cur = conn.cursor()

        cur.execute("""
            INSERT INTO afalpi (comentario, fecha, enviado_por)
            VALUES (%s, NOW(), %s)
        """, (comentario, session.get("usuario")))

        conn.commit()
        cur.close()
        conn.close()

        return jsonify({"message": "Pedido AFALPI guardado"}), 200

    except Exception as e:
        print("ERROR AFALPI:", e)
        return jsonify({"error": "Error interno"}), 500

# ================================
# GUARDAR PEDIDO AFALPI
# ================================
@app.route("/api/afalpi", methods=["POST"])
def guardar_afalpi():
    if "usuario" not in session:
        return jsonify({"error": "No autorizado"}), 401

    data = request.json
    comentario = data.get("comentario")

    if not comentario:
        return jsonify({"error": "Comentario vacío"}), 400

    conn = get_connection()
    if not conn:
        return jsonify({"error": "Error de conexión"}), 500

    cur = conn.cursor()
    try:
        cur.execute("""
            INSERT INTO afalpi (comentario, fecha, enviado_por)
            VALUES (%s, NOW(), %s)
        """, (comentario, session["usuario"]))

        conn.commit()
        return jsonify({"mensaje": "Pedido AFALPI enviado correctamente"})
    except Exception as e:
        conn.rollback()
        print("❌ Error AFALPI:", e)
        return jsonify({"error": "Error al guardar"}), 500
    finally:
        cur.close()
        conn.close()

# ================================
# LISTAR SOLICITUDES AFALPI
# ================================
@app.route("/api/afalpi", methods=["GET"])
@role_required("administrador")
def listar_afalpi():
    conn = get_connection()
    if not conn:
        return jsonify([]), 500

    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute("""
            SELECT id, comentario, fecha, enviado_por
            FROM afalpi
            ORDER BY fecha DESC
        """)
        data = cur.fetchall()
        return jsonify(data)
    except Exception as e:
        print("❌ Error listar AFALPI:", e)
        return jsonify([]), 500
    finally:
        cur.close()
        conn.close()

# ================================
# CREAR SOLICITUDES INVENTARIO
# ================================
@app.route("/api/inventario", methods=["POST"])
def crear_inventario():
    filas = request.json  # array de filas del formulario

    if not filas or len(filas) == 0:
        return jsonify({"error": "No hay datos"}), 400

    conn = get_connection()

    cur = conn.cursor()

    for r in filas:
        cur.execute("""
            INSERT INTO inventario
            (producto, cantidad, unidad, proveedor, comentario, enviado_por)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (
            r.get("producto"),
            r.get("cantidad"),
            r.get("unidad"),
            r.get("proveedor"),
            r.get("comentario"),
            session.get("usuario") 
        ))

    conn.commit()
    cur.close()
    conn.close()

    return jsonify({"ok": True})

# ================================
# GUARDAR PEDIDO INVENTARIO
# ================================

@app.route("/api/inventario", methods=["POST"])
def guardar_inventario():
    if "usuario" not in session:
        return jsonify({"error": "No autorizado"}), 401

    filas = request.json

    if not filas or len(filas) == 0:
        return jsonify({"error": "Sin datos"}), 400

    conn = get_connection()
    if not conn:
        return jsonify({"error": "Error de conexión"}), 500

    cur = conn.cursor()
    try:
        for r in filas:
            cur.execute("""
                INSERT INTO inventario
                (producto, cantidad, unidad, proveedor, comentario, fecha, enviado_por)
                VALUES (%s, %s, %s, %s, %s, NOW(), %s)
            """, (
                r.get("producto"),
                r.get("cantidad"),
                r.get("unidad"),
                r.get("proveedor"),
                r.get("comentario"),
                session["usuario"]
            ))

        conn.commit()
        return jsonify({"mensaje": "Pedido de inventario enviado correctamente"})
    except Exception as e:
        conn.rollback()
        print("❌ Error INVENTARIO:", e)
        return jsonify({"error": "Error al guardar"}), 500
    finally:
        cur.close()
        conn.close()

# ================================
# LISTAR PEDIDO INVENTARIO
# ================================

@app.route("/api/inventario", methods=["GET"])
@role_required("administrador")
def listar_inventario():
    conn = get_connection()
    if not conn:
        return jsonify([]), 500

    cur = conn.cursor(cursor_factory=RealDictCursor)
    try:
        cur.execute("""
            SELECT id, producto, cantidad, unidad, proveedor,
                   comentario, fecha, enviado_por
            FROM inventario
            ORDER BY fecha DESC
        """)
        data = cur.fetchall()
        return jsonify(data)
    except Exception as e:
        print("❌ Error listar INVENTARIO:", e)
        return jsonify([]), 500
    finally:
        cur.close()
        conn.close()

# ================================
# CONDIRMAR ELIMINAR PRODUCTO
# ================================

@app.route("/api/inventario/<int:id>", methods=["DELETE"])
@role_required("administrador")
def eliminar_inventario(id):
    conn = get_connection()
    if not conn:
        return jsonify({"error": "Error de conexión"}), 500

    cur = conn.cursor()
    try:
        cur.execute("DELETE FROM inventario WHERE id = %s", (id,))
        conn.commit()
        return jsonify({"mensaje": "Solicitud eliminada"})
    except Exception as e:
        conn.rollback()
        print("❌ Error eliminar INVENTARIO:", e)
        return jsonify({"error": "Error al eliminar"}), 500
    finally:
        cur.close()
        conn.close()

@app.route("/subir_programacion", methods=["GET", "POST"])
def subir_programacion():
    if "rol" not in session or session["rol"] != "administrador":
        return redirect("/login")

    if request.method == "POST":
        archivo = request.files.get("archivo")
        if archivo:
            nombre = archivo.filename
            contenido = archivo.read()

            conn = get_connection()
            cur = conn.cursor()
            cur.execute(
                "INSERT INTO programacion (nombre_archivo, contenido) VALUES (%s, %s)",
                (nombre, psycopg2.Binary(contenido))
            )
            conn.commit()
            cur.close()
            conn.close()
            return redirect("/")  # Regresa al home
    return render_template("subir_programacion.html")


# ================================
# RUN
# ================================
if __name__ == "__main__":
    app.run(debug=True, port=8080)
